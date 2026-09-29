import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'vendor/pptx-template-parser'))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'services/parser'))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'vendor/presentation-builder'))
from bridge import summarize
from template_parser.extract import extract_template
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

class BridgeTest(unittest.TestCase):
    def test_reads_actual_pptx_styles(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            prs=Presentation()
            slide=prs.slides.add_slide(prs.slide_layouts[6])
            run=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text_frame.paragraphs[0].add_run()
            run.text='Integration fixture'
            run.font.name='DejaVu Sans'
            run.font.size=Pt(24)
            run.font.color.rgb=RGBColor.from_string('123456')
            prs.save(root/'test.pptx')
            summary=summarize(extract_template(root/'test.pptx',root/'package'))
            self.assertEqual(summary['slides'],1)
            self.assertIn('#123456',[x['hex'] for x in summary['palette']])
            self.assertIn('DejaVu Sans',[x['family'] for x in summary['fonts']])
            self.assertIn(24,[x['pt'] for x in summary['fontSizes']])
            self.assertIsNone(summary['grid'])

if __name__=='__main__':unittest.main()
