import json
import sys
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parents[1]
for name in ('services/parser', 'vendor/pptx-template-parser', 'vendor/presentation-builder'):
    sys.path.insert(0, str(ROOT/name))
from generation import Generations, extract_content
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
from template_parser.extract import extract_template, sha256
from template_parser.composition import compose_slide
from template_parser.batch import run_batch, atomic_json
from template_parser.api import APIError
from template_parser.schemas import SlideSemantics
from template_parser.exporter import export_package
from presentation_builder.audit import audit_pptx


def fixture(root):
    prs = Presentation()
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text='Шаблон заголовка'
    prs.save(root/'source.pptx')
    package=root/'package'
    manifest=extract_template(root/'source.pptx',package)
    (package/'previews').mkdir()
    for slide in manifest['slides']:
        Image.new('RGB',(100,60)).save(package/'previews'/f"{slide['id']}.png")
    return package


def semantics(slide, *args, **kwargs):
    leaves=[o for o in slide['objects'] if o['kind']!='group']
    own=next(o for o in leaves if o['source_part'].lstrip('/')==slide['source_part'].lstrip('/') and o.get('text'))
    return SlideSemantics(slide_id=slide['id'],slide_role='title',usage='generation_template',needs_review=False,
        components=[{'id':'title','type':'title','members':[o['id'] for o in leaves],
                     'preserve':[o['id'] for o in leaves if o!=own],
                     'slots':[{'object_id':own['id'],'role':'title','content_type':'text'}]}],unassigned=[])


class GenerationTest(unittest.TestCase):
    def test_parser_export_to_three_editable_decks_and_resume(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);package=fixture(root)
            with patch('template_parser.batch.analyze_slide',side_effect=semantics):
                run_batch(package,{'model':'test'},delay=0)
            version=root/'version';export_package(package,version)
            jobs=SimpleNamespace(root=root/'jobs',executor=Mock(),capacity=threading.BoundedSemaphore(4),version=lambda *_:version)
            inputs=root/'inputs';inputs.mkdir();name='a'*32
            with zipfile.ZipFile(inputs/name,'w') as z:z.writestr('текст.md','Продукт помогает командам создавать презентации.')
            gen=Generations(jobs,root/'generations',inputs)
            data=dict(job_id='b'*32,parser_job_id='c'*32,version='d'*32,archive_sha256=sha256(version.with_suffix('.zip').read_bytes()),content_name=name,brief='На русском',slide_count=3)
            gen.submit(data);gen.submit(data);self.assertEqual(jobs.executor.submit.call_count,1)
            def plan(*args,**kwargs):
                return {'title':'Тест', 'slides':[{'template_slide_id':'s1','values':{'title.title':f'Новый текст {i}'},'source_refs':['asset-1']} for i in range(3)]}
            with patch('generation.LlmPlanner.plan',side_effect=plan) as calls:
                gen.process(data['job_id']);self.assertEqual(gen.state(data['job_id'])['status'],'done', (gen.folder(data['job_id'])/'error.json').read_text() if (gen.folder(data['job_id'])/'error.json').exists() else '')
                self.assertEqual(calls.call_count,3)
            for v in ('1','2','3'):
                path=gen.download(data['job_id'],v);self.assertTrue(audit_pptx(path)['ok'])
                pptx=Presentation(path);self.assertEqual(len(pptx.slides),3)
                self.assertIn('Новый текст',pptx.slides[0].shapes[0].text)
            gen.update(data['job_id'],status='interrupted');gen.resume(data['job_id'])
            with patch('generation.LlmPlanner.plan',side_effect=AssertionError('Must reuse saved plans')):
                gen.process(data['job_id']);self.assertEqual(gen.state(data['job_id'])['status'],'done')

    def test_transient_failure_does_not_skip_remaining_slides(self):
        with tempfile.TemporaryDirectory() as tmp:
            package=fixture(Path(tmp));index=0
            def answer(slide,*args,**kwargs):
                if slide['id']=='s2':raise APIError('category=timeout')
                return semantics(slide)
            with patch('template_parser.batch.analyze_slide',side_effect=answer):
                result=run_batch(package,{'model':'test','continue_transient_errors':True},delay=0)
            self.assertEqual(result['status'],'finished_with_errors')
            self.assertEqual(result['counts'],{'completed':2,'api_error':1})
            with patch('template_parser.batch.analyze_slide',side_effect=semantics) as calls:
                resumed=run_batch(package,{'model':'test','continue_transient_errors':True},delay=0)
            self.assertEqual(calls.call_count,1);self.assertEqual(resumed['counts'],{'skipped':2,'completed':1})

    def test_zip_traversal_and_invalid_job_paths_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            with zipfile.ZipFile(root/'bad.zip','w') as z:z.writestr('../escape.txt','x')
            with self.assertRaises(ValueError):extract_content(root/'bad.zip',root/'out')
            self.assertFalse((root/'escape.txt').exists())
            jobs=SimpleNamespace(root=root/'jobs')
            gen=Generations(jobs,root/'generations',root/'inputs')
            with self.assertRaises(FileNotFoundError):gen.state('../secrets')

if __name__=='__main__':unittest.main()
